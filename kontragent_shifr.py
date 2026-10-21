# -*- coding: utf-8 -*-
"""
Шифрование / расшифровка контрагентов в Excel по словарю.
Работает на Windows и macOS. Требуется: Python 3.9+ и openpyxl.

Словарь (dictionary.xlsx): 1-й столбец — исходное название, 2-й — зашифрованное,
первая строка — заголовок.
Файл данных: на каждом листе ищется столбец с заголовком «Контрагент» (1-я строка).
"""
from __future__ import annotations

import os
import platform
import subprocess
import threading
from pathlib import Path

from openpyxl import load_workbook

APP_TITLE = "Шифрование контрагентов"


# ---------------------------------------------------------------- логика ----
def _norm(v):
    """Сравниваем без лишних пробелов по краям."""
    return v.strip() if isinstance(v, str) else v


def load_dictionary(path: Path, decrypt: bool):
    """Возвращает (словарь замен, список предупреждений)."""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    mapping, warnings, seen = {}, [], {}
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if len(row) < 2:
            continue
        original, encrypted = _norm(row[0]), _norm(row[1])
        if original in (None, "") or encrypted in (None, ""):
            continue
        key, val = (encrypted, original) if decrypt else (original, encrypted)
        if key in seen and mapping[key] != val:
            warnings.append(
                f"Словарь, строки {seen[key]} и {i}: «{key}» встречается дважды "
                f"с разными значениями — используется строка {seen[key]}."
            )
            continue
        seen.setdefault(key, i)
        mapping.setdefault(key, val)
    wb.close()
    if not mapping:
        raise ValueError("Словарь пуст или в нём нет двух заполненных столбцов.")
    return mapping, warnings


def process(dict_path: Path, data_path: Path, out_path: Path,
            decrypt: bool, column: str = "Контрагент"):
    """Выполняет замену. Возвращает текстовый отчёт."""
    mapping, warnings = load_dictionary(dict_path, decrypt)
    wb = load_workbook(data_path)
    column = column.strip()
    total, sheets_found, missing = 0, 0, set()

    for ws in wb.worksheets:
        col = None
        for cell in ws[1]:
            if _norm(cell.value) == column:
                col = cell.column
                break
        if col is None:
            continue
        sheets_found += 1
        replaced = 0
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=col)
            v = _norm(cell.value)
            if v in (None, ""):
                continue
            if v in mapping:
                cell.value = mapping[v]
                replaced += 1
            else:
                missing.add(str(v))
        total += replaced
        warnings.append(f"Лист «{ws.title}»: заменено {replaced}")

    if sheets_found == 0:
        raise ValueError(f"Ни на одном листе не найден столбец «{column}» в первой строке.")

    wb.save(out_path)

    action = "Расшифровка" if decrypt else "Шифровка"
    lines = [f"{action} завершена. Всего заменено: {total}", *warnings]
    if missing:
        sample = sorted(missing)[:15]
        lines.append(f"Нет в словаре ({len(missing)} шт., оставлены как есть):")
        lines += [f"  • {m}" for m in sample]
        if len(missing) > 15:
            lines.append(f"  … и ещё {len(missing) - 15}")
    lines.append(f"Сохранено: {out_path}")
    return "\n".join(lines)


def open_folder(path: Path):
    folder = str(path.parent)
    try:
        if platform.system() == "Windows":
            os.startfile(folder)  # type: ignore[attr-defined]
        elif platform.system() == "Darwin":
            subprocess.Popen(["open", folder])
        else:
            subprocess.Popen(["xdg-open", folder])
    except Exception:
        pass


# --------------------------------------------------------------- интерфейс --
def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title(APP_TITLE)
    root.minsize(640, 460)

    dict_var = tk.StringVar()
    data_var = tk.StringVar()
    col_var = tk.StringVar(value="Контрагент")
    copy_var = tk.BooleanVar(value=True)
    last_out = {"path": None}

    xlsx_types = [("Excel", "*.xlsx *.xlsm"), ("Все файлы", "*.*")]

    def pick(var, title):
        p = filedialog.askopenfilename(title=title, filetypes=xlsx_types)
        if p:
            var.set(p)

    frm = ttk.Frame(root, padding=14)
    frm.pack(fill="both", expand=True)
    frm.columnconfigure(1, weight=1)

    ttk.Label(frm, text="Словарь:").grid(row=0, column=0, sticky="w", pady=4)
    ttk.Entry(frm, textvariable=dict_var).grid(row=0, column=1, sticky="ew", padx=6)
    ttk.Button(frm, text="Выбрать…",
               command=lambda: pick(dict_var, "Выберите словарь")).grid(row=0, column=2)

    ttk.Label(frm, text="Файл данных:").grid(row=1, column=0, sticky="w", pady=4)
    ttk.Entry(frm, textvariable=data_var).grid(row=1, column=1, sticky="ew", padx=6)
    ttk.Button(frm, text="Выбрать…",
               command=lambda: pick(data_var, "Выберите файл данных")).grid(row=1, column=2)

    ttk.Label(frm, text="Столбец:").grid(row=2, column=0, sticky="w", pady=4)
    ttk.Entry(frm, textvariable=col_var, width=24).grid(row=2, column=1, sticky="w", padx=6)

    ttk.Checkbutton(
        frm, variable=copy_var,
        text="Сохранять результат в новый файл (исходный не меняется)",
    ).grid(row=3, column=0, columnspan=3, sticky="w", pady=(4, 10))

    btns = ttk.Frame(frm)
    btns.grid(row=4, column=0, columnspan=3, sticky="ew")
    btn_enc = ttk.Button(btns, text="🔒  Зашифровать")
    btn_dec = ttk.Button(btns, text="🔓  Расшифровать")
    btn_open = ttk.Button(btns, text="Открыть папку", state="disabled",
                          command=lambda: last_out["path"] and open_folder(last_out["path"]))
    btn_enc.pack(side="left", expand=True, fill="x", padx=(0, 4))
    btn_dec.pack(side="left", expand=True, fill="x", padx=4)
    btn_open.pack(side="left", padx=(4, 0))

    log = tk.Text(frm, height=14, wrap="word", state="disabled")
    log.grid(row=5, column=0, columnspan=3, sticky="nsew", pady=(12, 0))
    frm.rowconfigure(5, weight=1)

    def write_log(text):
        log.configure(state="normal")
        log.delete("1.0", "end")
        log.insert("end", text)
        log.configure(state="disabled")

    def set_busy(busy):
        state = "disabled" if busy else "normal"
        btn_enc.configure(state=state)
        btn_dec.configure(state=state)
        root.configure(cursor="watch" if busy else "")

    def start(decrypt):
        d, f = Path(dict_var.get().strip()), Path(data_var.get().strip())
        if not dict_var.get().strip() or not d.is_file():
            messagebox.showwarning(APP_TITLE, "Выберите файл словаря.")
            return
        if not data_var.get().strip() or not f.is_file():
            messagebox.showwarning(APP_TITLE, "Выберите файл данных.")
            return
        if copy_var.get():
            suffix = "_расшифровано" if decrypt else "_зашифровано"
            out = f.with_name(f.stem + suffix + f.suffix)
        else:
            if not messagebox.askyesno(APP_TITLE, "Исходный файл будет перезаписан. Продолжить?"):
                return
            out = f

        set_busy(True)
        write_log("Обработка…")

        def worker():
            try:
                report = process(d, f, out, decrypt, col_var.get() or "Контрагент")
                ok = True
            except PermissionError:
                report = ("Не удалось сохранить файл — скорее всего, он открыт в Excel.\n"
                          "Закройте его и повторите.")
                ok = False
            except Exception as e:
                report = f"Ошибка: {e}"
                ok = False

            def done():
                set_busy(False)
                write_log(report)
                if ok:
                    last_out["path"] = out
                    btn_open.configure(state="normal")
                else:
                    messagebox.showerror(APP_TITLE, report)
            root.after(0, done)

        threading.Thread(target=worker, daemon=True).start()

    btn_enc.configure(command=lambda: start(False))
    btn_dec.configure(command=lambda: start(True))

    write_log("1. Выберите словарь (столбец A — исходное название, B — шифр).\n"
              "2. Выберите файл с данными (столбец «Контрагент»).\n"
              "3. Нажмите «Зашифровать» или «Расшифровать».")
    root.mainloop()


if __name__ == "__main__":
    run_gui()
